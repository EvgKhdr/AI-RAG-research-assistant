"""Load a stored document, split it into chunks, embed and index it."""
import re
from pathlib import Path

from sqlalchemy.orm import Session

import embeddings
import vectorstore
from config import CHUNK_OVERLAP, CHUNK_SIZE
from db import DOCS_DIR, Document, engine


def load_pages(path: Path) -> list[tuple[int | None, str]]:
    """Return [(page number or None, text)]."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        return [(i, p.extract_text() or "") for i, p in enumerate(PdfReader(path).pages, 1)]
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        parts += [c.text for t in d.tables for row in t.rows for c in row.cells]
        return [(None, "\n\n".join(parts))]
    return [(None, path.read_text(encoding="utf-8", errors="replace"))]  # txt, md


def _split_long(text: str, size: int) -> list[str]:
    """Split an oversized paragraph on sentence ends, hard-cutting anything still too long."""
    out: list[str] = []
    for sentence in re.split(r"(?<=[.!?])\s+", text):
        while len(sentence) > size:
            out.append(sentence[:size])
            sentence = sentence[size:]
        if sentence:
            out.append(sentence)
    return out


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    pieces: list[str] = []
    for para in re.split(r"\n\s*\n", text):
        para = re.sub(r"\s+", " ", para).strip()
        if para:
            pieces.extend([para] if len(para) <= size else _split_long(para, size))

    chunks: list[str] = []
    current = ""
    for piece in pieces:
        if current and len(current) + len(piece) + 2 > size:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            tail = tail.split(" ", 1)[-1] if " " in tail else tail  # start on a word boundary
            current = f"{tail}\n\n{piece}" if tail and len(tail) + len(piece) + 2 <= size else piece
        else:
            current = f"{current}\n\n{piece}" if current else piece
    if current:
        chunks.append(current)
    return chunks


def index_document(doc_id: int) -> int:
    """Index one stored document; returns the chunk count. Sets status to indexed/failed."""
    with Session(engine) as s:
        doc = s.get(Document, doc_id)
        if doc is None:
            return 0
        folder_id, filename = doc.folder_id, doc.filename
        path = DOCS_DIR / str(folder_id) / filename
        try:
            chunks = [
                {"text": t, "page": page}
                for page, text in load_pages(path)
                for t in chunk_text(text)
            ]
            if not chunks:
                raise ValueError("No extractable text (scanned PDF?)")
            vectors = embeddings.embed([c["text"] for c in chunks])
            vectorstore.add_chunks(doc_id, folder_id, filename, chunks, vectors)
            doc.status = "indexed"
            return len(chunks)
        except Exception:
            doc.status = "failed"
            raise
        finally:
            s.commit()
